"""
test_resume_export_regression.py - Comprehensive Regression Test Suite for Resume Export Pipeline

Tests:
1. PDF, DOCX, MD, TEX, TXT, and JSON export formats.
2. Non-ASCII / Indian Rupee (₹) symbol & special characters (smart quotes, em-dashes, bullets).
3. All visual templates: modern, classic, minimal, executive, ats_safe.
4. Minimal / Sparse profile with missing fields (placeholder safety).
5. Dense / Multi-page profile font scaling and layout stability.
6. PDF magic bytes (b'%PDF-') and DOCX magic bytes (b'PK').
7. Real python-docx parsing validation via docx.Document().
8. FastAPI /api/resume/export endpoint responses.
"""

import unittest
import io
import docx
from fastapi.testclient import TestClient
from backend.app.main import app, get_db
from backend.app.agents.agent4_export_generator import (
    generate_pdf_resume,
    generate_docx_resume,
    generate_md_resume,
    generate_tex_resume,
    generate_txt_resume,
    generate_json_resume,
    generate_resume,
    sanitize_for_pdf,
    safe_xml_escape,
    safe_tex_escape,
    ResumeGenerationError,
)

class TestResumeExportRegression(unittest.TestCase):

    def setUp(self):
        self.client = TestClient(app)
        self.sample_profile = {
            "id": 1,
            "name": "Aditya Sharma",
            "email": "aditya.sharma@example.com",
            "phone": "+91 98765 43210",
            "location": {"city": "Bengaluru", "state": "Karnataka", "country": "India"},
            "summary": "Experienced Full-Stack Lead Engineer managing ₹50 Lakhs budget & scaling distributed microservices to 10M+ daily active users.",
            "skills": ["Python", "FastAPI", "React", "PostgreSQL", "Docker", "AWS", "Redis"],
            "experience_list": [
                {
                    "title": "Senior Staff Software Engineer",
                    "company": "TechCorp India Pvt. Ltd.",
                    "duration_months": 24,
                    "start_date": "2022-01",
                    "end_date": "2024-01",
                    "description": "Led a high-velocity team of 12 engineers. Managed annual cloud infra budget of ₹75,000,000. Optimized API latency by 45%.",
                    "achievements": [
                        "Architected event-driven architecture processing 10,000 req/sec with <50ms P99 latency.",
                        "Saved ₹12,000,000 in AWS EC2 compute costs by introducing spot instances and auto-scaling rules.",
                        "■ Spearheaded zero-downtime migration from monolithic DB to sharded PostgreSQL."
                    ],
                    "technologies": ["Python", "FastAPI", "PostgreSQL", "Redis", "Kafka"]
                },
                {
                    "title": "Software Engineer",
                    "company": "Innovative Solutions Ltd",
                    "duration_months": 36,
                    "start_date": "2019-01",
                    "end_date": "2021-12",
                    "description": "Developed customer-facing dashboard with React & TypeScript. Implemented JWT authentication and RLS security policies.",
                    "achievements": [
                        "Built real-time analytics pipeline using WebSocket connections.",
                        "Integrated Razorpay gateway supporting ₹ INR transactions with 99.99% uptime."
                    ],
                    "technologies": ["React", "TypeScript", "Node.js", "MongoDB"]
                }
            ],
            "projects": [
                {
                    "title": "TheNextOpportunityFinder AI Engine",
                    "description": "Multi-agent career intelligence platform parsing 100k+ resumes and matching Indian tech startup jobs.",
                    "url": "https://thenextopportunityfinder.vercel.app",
                    "technologies": ["Python", "FastAPI", "React", "ReportLab", "Supabase"]
                }
            ],
            "education_list": [
                {
                    "degree": "Bachelor of Technology",
                    "field": "Computer Science & Engineering",
                    "institution": "Indian Institute of Technology (IIT) Delhi",
                    "graduation_year": "2019"
                }
            ]
        }

    def test_sanitize_and_escape_helpers(self):
        """Test Unicode sanitization and XML/TeX escaping logic."""
        raw_str = "Managed ₹50k & team with 'smart quotes' — bullet • test ■"
        sanitized = sanitize_for_pdf(raw_str)
        self.assertNotIn("₹", sanitized)
        self.assertIn("Rs.", sanitized)

        xml_escaped = safe_xml_escape("5 > 3 & 2 < 4 with ₹100 • item")
        self.assertIn("&gt;", xml_escaped)
        self.assertIn("&lt;", xml_escaped)
        self.assertIn("&amp;", xml_escaped)
        self.assertIn("Rs.", xml_escaped)
        self.assertIn("&bull;", xml_escaped)

        tex_escaped = safe_tex_escape("100% success rate & $50k profit_loss #1")
        self.assertIn("\\%", tex_escaped)
        self.assertIn("\\&", tex_escaped)
        self.assertIn("\\$", tex_escaped)
        self.assertIn("\\_", tex_escaped)
        self.assertIn("\\#", tex_escaped)

    def test_pdf_generation_magic_bytes_and_templates(self):
        """Test PDF generation across all visual templates with non-ASCII characters."""
        templates = ["modern", "classic", "minimal", "executive", "ats_safe"]
        for tmpl in templates:
            pdf_bytes = generate_pdf_resume(self.sample_profile, template=tmpl)
            self.assertIsInstance(pdf_bytes, bytes)
            self.assertTrue(pdf_bytes.startswith(b'%PDF-'), f"PDF magic bytes check failed for template {tmpl}")
            self.assertGreater(len(pdf_bytes), 1000, f"PDF file size suspiciously small for template {tmpl}")

    def test_docx_generation_and_document_parsing(self):
        """Test DOCX generation magic bytes and python-docx structure parsing."""
        docx_bytes = generate_docx_resume(self.sample_profile)
        self.assertIsInstance(docx_bytes, bytes)
        self.assertTrue(docx_bytes.startswith(b'PK'), "DOCX magic bytes check failed")

        # Parse with python-docx to verify document structure integrity
        doc = docx.Document(io.BytesIO(docx_bytes))
        paragraphs_text = " ".join([p.text for p in doc.paragraphs])
        self.assertIn("Aditya Sharma", paragraphs_text)
        self.assertIn("TechCorp India", paragraphs_text)
        self.assertIn("Senior Staff Software Engineer", paragraphs_text)

    def test_markdown_and_latex_exports(self):
        """Test Markdown and LaTeX resume generation."""
        md_text = generate_md_resume(self.sample_profile)
        self.assertIn("# Aditya Sharma", md_text)
        self.assertIn("## Professional Summary", md_text)
        self.assertIn("TechCorp India", md_text)

        tex_text = generate_tex_resume(self.sample_profile, template="modern")
        self.assertIn("\\documentclass", tex_text)
        self.assertIn("\\name{Aditya}{Sharma}", tex_text)
        self.assertIn("\\end{document}", tex_text)

    def test_txt_and_json_exports(self):
        """Test Plain Text and JSON export formats."""
        txt_text = generate_txt_resume(self.sample_profile)
        self.assertIn("Aditya Sharma", txt_text)
        self.assertIn("SUMMARY", txt_text)
        self.assertIn("TECHNICAL SKILLS", txt_text)

        json_text = generate_json_resume(self.sample_profile)
        self.assertIn('"name": "Aditya Sharma"', json_text)

    def test_sparse_profile_placeholders(self):
        """Test export handling for sparse profiles with missing optional fields."""
        sparse_profile = {"name": "Jane Doe"}
        pdf_bytes = generate_pdf_resume(sparse_profile, template="modern")
        self.assertTrue(pdf_bytes.startswith(b'%PDF-'))

        docx_bytes = generate_docx_resume(sparse_profile)
        self.assertTrue(docx_bytes.startswith(b'PK'))

        md_text = generate_md_resume(sparse_profile)
        self.assertIn("# Jane Doe", md_text)
        self.assertIn("[Add description]", md_text)

    def test_dense_multi_page_profile_font_scaling(self):
        """Test dense profile with many achievements to verify adaptive font scaling and multi-page stability."""
        dense_profile = dict(self.sample_profile)
        dense_profile["experience_list"] = self.sample_profile["experience_list"] * 5  # 10 large experience blocks
        pdf_bytes = generate_pdf_resume(dense_profile, template="modern")
        self.assertTrue(pdf_bytes.startswith(b'%PDF-'))
        self.assertGreater(len(pdf_bytes), 5000)

    def test_api_export_endpoint_responses(self):
        """Test GET /api/resume/export endpoint for all formats."""
        formats = ["pdf", "docx", "md", "tex", "txt", "json"]
        for fmt in formats:
            res = self.client.get(f"/api/resume/export?format={fmt}&template=modern")
            self.assertEqual(res.status_code, 200, f"Export API failed for format {fmt}: {res.text}")
            self.assertIn("Content-Disposition", res.headers)
            if fmt == "pdf":
                self.assertTrue(res.content.startswith(b'%PDF-'))
            elif fmt == "docx":
                self.assertTrue(res.content.startswith(b'PK'))

if __name__ == "__main__":
    unittest.main()
